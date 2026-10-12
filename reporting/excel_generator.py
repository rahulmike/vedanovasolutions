"""Executive Excel Workbook Generator for NexaData AI."""
import io
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_executive_excel(
    question: str,
    query: str,
    df_result: pd.DataFrame,
    dataset_name: str = "datainfo.csv"
) -> bytes:
    """
    Generate an executive-styled multi-tab Excel workbook with formatted headers,
    metadata sheet, and auto-fitted column widths.
    """
    buffer = io.BytesIO()
    
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Sheet 1: Results
        df_result.to_excel(writer, sheet_name="Analytical Findings", index=False)
        
        # Sheet 2: Query & Metadata
        meta_df = pd.DataFrame([
            {"Parameter": "Application", "Value": "NexaData AI"},
            {"Parameter": "Generated Date", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"Parameter": "Source Dataset", "Value": dataset_name},
            {"Parameter": "Analytical Question", "Value": question},
            {"Parameter": "SQL Query Executed", "Value": query},
            {"Parameter": "Total Records", "Value": len(df_result)}
        ])
        meta_df.to_excel(writer, sheet_name="Report Metadata", index=False)
        
        workbook = writer.book
        
        # --- Style Sheet 1: Analytical Findings ---
        ws_data = workbook["Analytical Findings"]
        
        # Header Style: Navy blue with white bold text
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        
        # Apply header styling
        for col_idx in range(1, len(df_result.columns) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Apply alternating row colors and borders
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        for row_idx in range(2, len(df_result) + 2):
            for col_idx in range(1, len(df_result.columns) + 1):
                cell = ws_data.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if row_idx % 2 == 1:
                    cell.fill = alt_fill

        # Auto-adjust column widths
        for col in ws_data.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
            
        # --- Style Sheet 2: Metadata ---
        ws_meta = workbook["Report Metadata"]
        for col_idx in (1, 2):
            cell = ws_meta.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
        for col in ws_meta.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_meta.column_dimensions[col_letter].width = min(max(max_len + 3, 15), 60)

    buffer.seek(0)
    return buffer.getvalue()
